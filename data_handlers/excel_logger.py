# ============ data_handlers/excel_logger.py ============
"""
Καταγραφή σε Excel με openpyxl και χρωματική επισήμανση γραμμών.

    ΚΟΚΚΙΝΟ  → ερώτηση για αλλεργιογόνα
    ΚΙΤΡΙΝΟ  → ερώτηση για τιμή
    ΠΡΑΣΙΝΟ  → αίτημα κράτησης
    ΜΠΛΕ     → επιβεβαίωση παραγγελίας
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def make_fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


class ExcelSheetWriter:
    """Γενικός writer: ένα αρχείο, ένα φύλλο με επικεφαλίδες, προσθήκη γραμμών."""

    def __init__(self, path: Path, sheet_name: str, headers: list[str], widths: list[int] | None = None):
        self.path = Path(path)
        self.sheet_name = sheet_name
        self.headers = headers
        self.widths = widths or [18] * len(headers)

    def _open(self):
        if self.path.exists():
            wb = load_workbook(self.path)
            if self.sheet_name in wb.sheetnames:
                return wb, wb[self.sheet_name]
            ws = wb.create_sheet(self.sheet_name)
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = self.sheet_name
        self._write_header(ws)
        return wb, ws

    def _write_header(self, ws):
        for col, (title, width) in enumerate(zip(self.headers, self.widths), start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

    def _save(self, wb) -> Path:
        try:
            wb.save(self.path)
            return self.path
        except PermissionError:
            # Συνήθως στα Windows: το αρχείο είναι ανοιχτό στο Excel
            backup = self.path.with_name(f"{self.path.stem}_{datetime.now():%Y%m%d_%H%M%S}{self.path.suffix}")
            wb.save(backup)
            msg = f"Το {self.path.name} είναι ανοιχτό (π.χ. στο Excel). Αποθηκεύτηκε αντίγραφο: {backup.name}"
            log.warning(msg)
            print(f"  ⚠ {msg}")
            return backup

    def append_rows(self, rows: list[list], fill_color: str | None = None) -> Path:
        wb, ws = self._open()
        fill = make_fill(fill_color) if fill_color else None
        for values in rows:
            ws.append(values)
            r = ws.max_row
            for col in range(1, len(self.headers) + 1):
                cell = ws.cell(row=r, column=col)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if fill:
                    cell.fill = fill
        return self._save(wb)

    def read_rows(self) -> list[dict]:
        if not self.path.exists():
            return []
        wb = load_workbook(self.path, read_only=True)
        if self.sheet_name not in wb.sheetnames:
            return []
        rows = list(wb[self.sheet_name].iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        header = list(rows[0])
        return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


# ---------------------------------------------------------------------------
class ChatLogger:
    """chat_log.xlsx: κάθε μήνυμα πελάτη + απάντηση + λόγος επισήμανσης."""

    HEADERS = ["Timestamp", "Session ID", "User Message", "Detected Intent (classifier)", "LLM Intent",
               "Action", "Bot Response", "Highlight Reason"]
    WIDTHS = [20, 14, 45, 24, 18, 18, 70, 22]

    def __init__(self, path: Path = config.CHAT_LOG_XLSX):
        self.writer = ExcelSheetWriter(path, "Chat Log", self.HEADERS, self.WIDTHS)
        self._ensure_legend()

    def _ensure_legend(self):
        path = self.writer.path
        if path.exists():
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Chat Log"
        self.writer._write_header(ws)
        legend = wb.create_sheet("Legend")
        legend.append(["Highlight Reason", "Colour", "Meaning"])
        meanings = {
            "Allergen question": "Κόκκινο – ερώτηση για αλλεργιογόνα",
            "Price question": "Κίτρινο – ερώτηση για τιμή",
            "Reservation request": "Πράσινο – αίτημα κράτησης",
            "Order confirmation": "Μπλε – επιβεβαίωση παραγγελίας",
        }
        for reason, color in config.HIGHLIGHT_COLORS.items():
            legend.append([reason, color, meanings[reason]])
            for c in legend[legend.max_row]:
                c.fill = make_fill(color)
        legend.append(["(empty)", "—", "Χωρίς επισήμανση: γενικές ερωτήσεις / small talk"])
        for col, w in zip("ABC", (22, 10, 45)):
            legend.column_dimensions[col].width = w
        for c in legend[1]:
            c.font = Font(bold=True)
        self.writer._save(wb)

    def log(self, session_id: str, user_message: str, classifier_intent: str, llm_intent: str,
            action: str, response: str, highlight_reason: str | None) -> None:
        color = config.HIGHLIGHT_COLORS.get(highlight_reason) if highlight_reason else None
        self.writer.append_rows([[
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"), session_id, user_message, classifier_intent,
            llm_intent, action, response, highlight_reason or "",
        ]], fill_color=color)

# ============ training/evaluate.py ============
"""
Αξιολόγηση για την αναφορά της εργασίας.

Μετρικές:
  • JSON validity rate        – πόσες φορές το μοντέλο έβγαλε έγκυρο JSON
  • Intent accuracy + precision / recall / F1 ανά intent + confusion matrix
  • Action accuracy
  • Argument exact match      – σωστά πιάτα, ποσότητες, αλλεργιογόνα, ημερομηνία/ώρα/άτομα
  • Συμφωνία PyTorch classifier ↔ LLM
  • Full pipeline accuracy  – η τελική απόφαση ΜΕΤΑ τα guard rails (αυτό που βλέπει ο πελάτης)
  • (προαιρετικά) σύγκριση base μοντέλου (χωρίς fine-tuning) vs fine-tuned

Σύνολα δεδομένων: data/val.jsonl (συνθετικό) και data/test_dialogues.jsonl
(30 χειρόγραφα, ρεαλιστικά μηνύματα που ΔΕΝ προέρχονται από τα templates).

    python -m training.evaluate                  # fine-tuned LLM + classifier
    python -m training.evaluate --compare-base   # + base Qwen χωρίς LoRA (zero-shot)
    python -m training.evaluate --skip-llm       # μόνο ο classifier (δευτερόλεπτα)
    python -m training.evaluate --limit 100      # λιγότερα val δείγματα (γρηγορότερο σε CPU)
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config  # noqa: E402
from chatbot.intent_classifier import IntentPredictor  # noqa: E402
from chatbot.menu_data import Menu, normalize  # noqa: E402
from data_handlers.reservation_manager import TZ, parse_date, parse_time, to_int  # noqa: E402
from training.dataset_builder import load_jsonl  # noqa: E402

MENU = Menu()
NOW = datetime.now(TZ)


# ---------------------------------------------------------------------------
def load_examples(path: Path, limit: int = 0) -> list[dict]:
    """Ενιαία μορφή: message, history, cart, pending, expected{intent, action, arguments}."""
    out = []
    for r in load_jsonl(path):
        if "meta" in r:   # από dataset_builder
            out.append({"id": r["id"], "message": r["meta"]["message"], "history": r["meta"]["history"],
                        "cart": r["meta"]["cart"], "pending": r["meta"]["pending"], "expected": r["target"]})
        else:             # χειρόγραφα test dialogues
            out.append({"id": r["id"], "message": r["message"], "history": r.get("history", []),
                        "cart": r.get("cart", []), "pending": r.get("pending"), "expected": r["expected"]})
    return out[:limit] if limit else out


def canon_args(action: str, args: dict | None):
    """Κανονικοποίηση ορισμάτων ώστε π.χ. 'καρμπονάρα' == 'Spaghetti Carbonara'."""
    args = args or {}

    def dish_id(name):
        d = MENU.find_dish(name) if name else None
        return d["id"] if d else (normalize(name) if name else None)

    if action == "get_menu":
        return MENU.find_category(args.get("category")) if args.get("category") else None
    if action in ("get_dish_info", "get_price"):
        return dish_id(args.get("dish"))
    if action == "get_allergens":
        return dish_id(args.get("dish")), MENU.find_allergen(args.get("allergen")) if args.get("allergen") else None
    if action in ("add_item", "remove_item"):
        items = args.get("items") or []
        return tuple(sorted((dish_id(i.get("name")), i.get("qty")) for i in items if isinstance(i, dict)))
    if action == "make_reservation":
        return (to_int(args.get("party_size")), parse_date(args.get("date_text"), NOW),
                parse_time(args.get("time_text")), normalize(args.get("name") or "") or None)
    return None


def run_llm(llm, examples: list[dict], label: str) -> list[dict]:
    from tqdm import tqdm
    from chatbot.prompt_format import build_messages, build_user_content, parse_model_output

    preds = []
    for ex in tqdm(examples, desc=label):
        t0 = time.time()
        msgs = build_messages(build_user_content([tuple(h) for h in ex["history"]], ex["cart"], ex["pending"],
                                                 ex["message"]))
        raw = llm.generate_raw(msgs, greedy=True)
        preds.append({"parsed": parse_model_output(raw), "raw": raw, "latency": time.time() - t0})
    return preds


def pipeline_accuracy(examples: list[dict], parsed_preds: list[dict | None] | None = None) -> tuple[float, list]:
    """Ακρίβεια ΟΛΟΚΛΗΡΟΥ του συστήματος: (LLM ή κανόνες) + guard rails + επαλήθευση οντοτήτων.
    Δεν εκτελεί εργαλεία (δεν γράφει παραγγελίες/κρατήσεις), μόνο συγκρίνει την τελική απόφαση."""
    from chatbot.conversation import RestaurantChatbot
    from chatbot.rule_parser import rule_decide

    bot = RestaurantChatbot(use_llm=False)
    ok, wrong = 0, []
    for k, ex in enumerate(examples):
        s = bot.new_session()
        s.pending = ex["pending"]
        for h in ex["history"]:
            s.history.append(tuple(h))
        for it in ex["cart"]:
            d = MENU.find_dish(it["name"])
            if d:
                s.cart.add(d, it["qty"])
        cls, conf = bot.classifier.predict(ex["message"], s.pending)
        parsed = parsed_preds[k] if parsed_preds else None
        dec = parsed or rule_decide(ex["message"], cls, bot.menu, s.history, s.pending, bool(s.cart.items), s.cart.items)
        dec = bot._guard(dec, ex["message"], cls, conf, s)
        exp = ex["expected"]
        good = dec["action"] == exp["action"] and (
            exp["action"] == "none" or canon_args(dec["action"], dec["arguments"]) == canon_args(exp["action"], exp.get("arguments")))
        ok += good
        if not good:
            wrong.append((ex, dec))
    return ok / max(1, len(examples)), wrong


def score(examples, preds, classifier_preds=None) -> dict:
    n = len(examples)
    valid = intent_ok = action_ok = args_ok = agree = 0
    y_true, y_pred, errors = [], [], []
    for ex, p, cp in zip(examples, preds, classifier_preds or [None] * n):
        exp = ex["expected"]
        got = p["parsed"]
        y_true.append(exp["intent"])
        y_pred.append(got["intent"] if got else "invalid_json")
        if got is None:
            errors.append((ex, p, "invalid JSON"))
            continue
        valid += 1
        i_ok = got["intent"] == exp["intent"]
        a_ok = got["action"] == exp["action"]
        g_ok = a_ok and canon_args(got["action"], got["arguments"]) == canon_args(exp["action"], exp.get("arguments"))
        intent_ok += i_ok
        action_ok += a_ok
        args_ok += g_ok
        if cp is not None:
            agree += cp == got["intent"]
        if not (i_ok and a_ok and g_ok):
            errors.append((ex, p, "wrong intent" if not i_ok else ("wrong action" if not a_ok else "wrong arguments")))
    lat = [p.get("latency", 0) for p in preds]
    return {"n": n, "json_validity": valid / n, "intent_acc": intent_ok / n, "action_acc": action_ok / n,
            "args_exact_match": args_ok / n, "agreement_with_classifier": (agree / n) if classifier_preds else None,
            "avg_latency_s": sum(lat) / max(1, len(lat)), "y_true": y_true, "y_pred": y_pred, "errors": errors}


def classifier_scores(examples, cls_preds) -> dict:
    y_true = [ex["expected"]["intent"] for ex in examples]
    acc = sum(t == p for t, p in zip(y_true, cls_preds)) / len(examples)
    return {"n": len(examples), "intent_acc": acc, "y_true": y_true, "y_pred": cls_preds}


# ---------------------------------------------------------------------------
def plot_confusion(y_true, y_pred, title: str, path: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from sklearn.metrics import confusion_matrix

    labels = list(config.INTENTS) + (["invalid_json"] if "invalid_json" in y_pred else [])
    cm = confusion_matrix(y_true, y_pred, labels=labels)
    fig, ax = plt.subplots(figsize=(9, 7.5))
    im = ax.imshow(cm, cmap="Blues")
    ax.set_xticks(range(len(labels)), labels, rotation=45, ha="right", fontsize=8)
    ax.set_yticks(range(len(labels)), labels, fontsize=8)
    for i in range(len(labels)):
        for j in range(len(labels)):
            if cm[i, j]:
                ax.text(j, i, cm[i, j], ha="center", va="center", fontsize=8,
                        color="white" if cm[i, j] > cm.max() / 2 else "black")
    ax.set_xlabel("predicted")
    ax.set_ylabel("true")
    ax.set_title(title)
    fig.colorbar(im, ax=ax, fraction=0.046)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def write_report(summary_rows, per_intent_rows, error_rows) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    sheets = [("Summary", summary_rows), ("Per Intent", per_intent_rows), ("Errors", error_rows)]
    for idx, (name, rows) in enumerate(sheets):
        ws = wb.active if idx == 0 else wb.create_sheet(name)
        ws.title = name
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(12, *(len(str(c.value or "")) for c in col[:50])))
        ws.freeze_panes = "A2"
        if name == "Errors":
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    wb.save(config.EVAL_XLSX)


def main():
    ap = argparse.ArgumentParser(description="Αξιολόγηση classifier + fine-tuned LLM")
    ap.add_argument("--limit", type=int, default=0, help="πλήθος val δειγμάτων (0 = όλα)")
    ap.add_argument("--skip-llm", action="store_true")
    ap.add_argument("--compare-base", action="store_true", help="αξιολόγηση και του base μοντέλου χωρίς LoRA")
    ap.add_argument("--adapter", default=str(config.ADAPTER_DIR))
    ap.add_argument("--base-model", default=None)
    args = ap.parse_args()

    from sklearn.metrics import precision_recall_fscore_support

    datasets = {"validation (synthetic)": load_examples(config.VAL_PATH, args.limit),
                "test dialogues (hand-written)": load_examples(config.TEST_DIALOGUES_PATH)}
    clf = IntentPredictor()
    summary = [["Model", "Dataset", "N", "JSON validity", "Intent acc", "Action acc", "Args exact match",
                "Agreement w/ classifier", "Avg latency (s)", "Full pipeline acc (action+args)"]]
    per_intent = [["Model", "Dataset", "Intent", "Precision", "Recall", "F1", "Support"]]
    errors = [["Model", "Dataset", "Error", "Message", "Expected", "Predicted", "Raw output"]]

    def add_per_intent(model, ds, y_true, y_pred):
        p, r, f, s = precision_recall_fscore_support(y_true, y_pred, labels=config.INTENTS, zero_division=0)
        for i, lab in enumerate(config.INTENTS):
            if s[i]:
                per_intent.append([model, ds, lab, round(p[i], 3), round(r[i], 3), round(f[i], 3), int(s[i])])

    models = []
    if not args.skip_llm:
        from chatbot.local_llm import LocalLLM, adapter_available
        if not adapter_available(Path(args.adapter)):
            print("✖ Δεν βρέθηκε adapter. Τρέξτε πρώτα: python -m training.train  (ή --skip-llm)")
            sys.exit(1)
        models.append(("fine-tuned (LoRA)", lambda: LocalLLM(Path(args.adapter), args.base_model, use_adapter=True)))
        if args.compare_base:
            models.append(("base (zero-shot)", lambda: LocalLLM(Path(args.adapter), args.base_model, use_adapter=False)))

    confusion_done = False
    for ds_name, examples in datasets.items():
        cls_preds = [clf.predict(ex["message"], ex["pending"])[0] for ex in examples]
        cs = classifier_scores(examples, cls_preds)
        pipe_acc, _ = pipeline_accuracy(examples)
        summary.append(["PyTorch intent classifier (+ rules)", ds_name, cs["n"], "-", round(cs["intent_acc"], 3),
                        "-", "-", "-", "-", round(pipe_acc, 3)])
        add_per_intent("PyTorch intent classifier", ds_name, cs["y_true"], cs["y_pred"])
        print(f"[classifier] {ds_name}: intent acc = {cs['intent_acc']:.3f} | full pipeline (rules) = {pipe_acc:.3f}")
        if args.skip_llm and not confusion_done:
            plot_confusion(cs["y_true"], cs["y_pred"], f"PyTorch intent classifier – {ds_name}", config.CONFUSION_PNG)
            confusion_done = True

    for model_name, factory in models:
        llm = factory()
        for ds_name, examples in datasets.items():
            cls_preds = [clf.predict(ex["message"], ex["pending"])[0] for ex in examples]
            preds = run_llm(llm, examples, f"{model_name} · {ds_name}")
            sc = score(examples, preds, cls_preds)
            pipe_acc, _ = pipeline_accuracy(examples, [p["parsed"] for p in preds])
            summary.append([model_name, ds_name, sc["n"], round(sc["json_validity"], 3), round(sc["intent_acc"], 3),
                            round(sc["action_acc"], 3), round(sc["args_exact_match"], 3),
                            round(sc["agreement_with_classifier"], 3), round(sc["avg_latency_s"], 2), round(pipe_acc, 3)])
            add_per_intent(model_name, ds_name, sc["y_true"], sc["y_pred"])
            for ex, p, why in sc["errors"][:200]:
                errors.append([model_name, ds_name, why, ex["message"], json.dumps(ex["expected"], ensure_ascii=False),
                               json.dumps(p["parsed"], ensure_ascii=False) if p["parsed"] else "", p["raw"][:300]])
            print(f"[{model_name}] {ds_name}: JSON {sc['json_validity']:.3f} | intent {sc['intent_acc']:.3f} | "
                  f"action {sc['action_acc']:.3f} | args {sc['args_exact_match']:.3f} | full pipeline {pipe_acc:.3f}")
            if not confusion_done and model_name.startswith("fine-tuned"):
                plot_confusion(sc["y_true"], sc["y_pred"], f"Fine-tuned LLM – {ds_name}", config.CONFUSION_PNG)
                confusion_done = True
        del llm

    write_report(summary, per_intent, errors)
    print(f"\n✔ Αναφορά: {config.EVAL_XLSX}\n✔ Confusion matrix: {config.CONFUSION_PNG}")


if __name__ == "__main__":
    main()
